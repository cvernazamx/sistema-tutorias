import imaplib
import email
from email.header import decode_header
import re
from datetime import datetime, date
import pandas as pd
import socket
import ssl
import time
import unicodedata
import io
import logging

from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# Silenciar advertencias internas de pypdf
try:
from pypdf import PdfReader
PYPDF_DISPONIBLE = True
logging.getLogger("pypdf").setLevel(logging.ERROR)
except ImportError:
try:
from PyPDF2 import PdfReader
PYPDF_DISPONIBLE = True
logging.getLogger("PyPDF2").setLevel(logging.ERROR)
except ImportError:
PYPDF_DISPONIBLE = False

# ==========================================
# 1. CONFIGURACIÓN DE ACCESO
# ==========================================
IMAP_SERVER = "mail.tce.gob.ec"
IMAP_PORT = 993
USERNAME = "secretaria.general@tce.gob.ec" # Coloca tu usuario o correo institucional
PASSWORD = r"Sg.TcE.2025$" # Coloca tu contraseña entre r"...


INBOX_FOLDER = "INBOX"
SENT_FOLDER = "Sent"

# Rango evaluado: Desde el 10 de Agosto de 2026 hasta hoy
FECHA_DESDE_STR = "10-Aug-2026"
FECHA_AUDITORIA_AHORA = datetime.now()

socket.setdefaulttimeout(45)

print("=" * 65)
print(f" AUDITORÍA DE CERTIFICACIONES Y SOLICITUDES TCE")
print(f" Fecha y Hora de Ejecución: {FECHA_AUDITORIA_AHORA.strftime('%Y-%m-%d %H:%M:%S')}")
print(f" Período Evaluado: Desde {FECHA_DESDE_STR} hasta {FECHA_AUDITORIA_AHORA.strftime('%d-%b-%Y')}")
print("=" * 65)

# ==========================================
# 2. FUNCIONES DE NORMALIZACIÓN Y PARSEO
# ==========================================
def remove_accents(input_str):
if not input_str:
return ""
nfkd_form = unicodedata.normalize('NFKD', input_str)
return "".join([c for c in nfkd_form if not unicodedata.combining(c)])

def clean_header(header_val):
if not header_val:
return ""
decoded = decode_header(header_val)
parts = []
for text, enc in decoded:
if isinstance(text, bytes):
try:
parts.append(text.decode(enc or 'utf-8', errors='ignore'))
except Exception:
parts.append(text.decode('latin-1', errors='ignore'))
else:
parts.append(str(text))
return "".join(parts).strip()

def extract_email_address(raw_str):
match = re.search(r'[\w\.-]+@[\w\.-]+', raw_str)
return match.group(0).lower() if match else ""

def matches_subject_and_sender_criteria(subject, sender_raw):
sender_clean = remove_accents(sender_raw).lower()
sender_email = extract_email_address(sender_raw)
text_clean = remove_accents(subject).lower()
text_raw = subject.lower()
if "documentos@cne.gob.ec" in sender_email or "edoc" in sender_clean or "edoc" in text_clean:
return False
strict_patterns = [
# Recursos en contra de resoluciones al TCE
r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+(?:presentados?|interpuestos?)\s+(?:en\s+contra\s+de|contra)\s+(?:las\s+)?resoluci[oó]n(?:es)?',
r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+(?:presentados?|interpuestos?)\s+(?:en\s+contra\s+de|contra)',
r'(?:recursos?|resoluci[oó]n(?:es)?)\s+(?:en\s+contra\s+de|contra)\s+(?:las\s+)?resoluci[oó]n(?:es)?\s+(?:al|ante\s+el)\s+tce',

# Certificación de Recursos Subjetivos Contenciosos Electorales
r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+subjetivos?(?:\s+contenciosos?)?(?:\s+electorales?)?',
r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+subjetivos?',

# Certificación de recurso(s) pendiente(s)
r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+pendientes?',
r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+pendientes?',
# No tener recursos pendientes
r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?(?:no\s+)?(?:tener|poseer|registrar|existir)\s+recursos?\s+pendientes?',
r'no\s+tener\s+recursos?\s+pendientes?',
# Certificación de Resoluciones
r'(?:solicitud|solicito|peticion|pedido)\s+(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:las\s+)?resoluci[oó]n(?:es)?',
r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:las\s+)?resoluci[oó]n(?:es)?',
# Certificación de Recursos
r'(?:solicitud|solicito|peticion|pedido)\s+(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?',
r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?',
r'certificaci[oó]n\s+(?:de\s+)?no\s+(?:haber\s+)?interpuesto\s+recursos?',
# Certificaciones generales
r'certificaci[oó]n',
r'certificaciones',
r'solicitud\s+(?:de\s+)?certificaci[oó]n',
r'solicito\s+certificaci[oó]n',
r'certificacion',
r'certificación',
# Juntas Provinciales Electorales / JPEs / Delegaciones
r'junta[s]?\s+provincial(?:es)?',
r'junta[s]?\s+provincial(?:es)?\s+electoral(?:es)?',
r'\bjpe[a-z0-9_\-]*\b',
r'delegaci[oó]n\s+provincial',
# Interpuesto
r'interpuest[oa]s?',
r'interposici[oó]n',
r'interposicion',
r'interposición'
]
return any(re.search(pat, text_clean) or re.search(pat, text_raw) for pat in strict_patterns)

def parse_email_content_and_attachments(msg, leer_adjuntos_pdf=False):
body = ""
attachments = []
pdf_text_content = ""
if msg.is_multipart():
for part in msg.walk():
ctype = part.get_content_type()
cdispo = str(part.get('Content-Disposition', ''))
filename = part.get_filename()
if filename:
clean_fname = clean_header(filename)
fname_lower = clean_fname.lower()
if not (fname_lower.endswith('.doc') or fname_lower.endswith('.docx')):
attachments.append(clean_fname)
if leer_adjuntos_pdf and fname_lower.endswith('.pdf') and PYPDF_DISPONIBLE:
try:
payload = part.get_payload(decode=True)
if payload:
pdf_file = io.BytesIO(payload)
reader = PdfReader(pdf_file, strict=False)
for page_idx in range(min(2, len(reader.pages))):
try:
extracted = reader.pages[page_idx].extract_text()
if extracted:
pdf_text_content += extracted + " "
except Exception:
pass
except Exception:
pass
elif 'attachment' in cdispo:
attachments.append("adjunto_sin_nombre")
if ctype in ['text/plain', 'text/html'] and 'attachment' not in cdispo:
try:
payload = part.get_payload(decode=True)
if payload:
decoded = payload.decode('utf-8', errors='ignore')
decoded = re.sub(r'<[^>]+>', ' ', decoded)
body += decoded + " "
except Exception:
pass
else:
try:
payload = msg.get_payload(decode=True)
if payload:
body = payload.decode('utf-8', errors='ignore')
except Exception:
pass

full_body = (body + " " + pdf_text_content).strip()
return full_body, attachments

# ==========================================================
# 3. DETECCIÓN FLEXIBLE Y EXACTA DE OFICIOS (ENTRADA Y SALIDA)
# ==========================================================
PATRONES_OFICIO_TCE = [
r'\b(TCE-(?:SG|P|OM)(?:-[A-Za-z0-9]+)*-[0-9]{4}-[0-9]{1,5}-[O0o])\b',
r'\b(TCE-[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*-[0-9]{4}-[0-9]{1,5}(?:-[O0o])?)\b',
r'(?:oficio|oficio\s+nro\.?|oficio\s+n°|oficio\s+num\.?)\s*[:#-]?\s*([0-9]{1,5}-[0-9]{4}-TCE(?:-[A-Za-z0-9]+)?)',
r'(?:oficio|oficio\s+nro\.?|oficio\s+n°|oficio\s+num\.?)\s*[:#-]?\s*(TCE-[A-Za-z0-9_\-\/]+)'
]

def extract_oficio_tce_exacto(text):
if not text:
return None
for pat in PATRONES_OFICIO_TCE:
match = re.search(pat, text, re.IGNORECASE)
if match:
val = match.group(1).strip().strip('.,;')
if val.endswith('-0'):
val = val[:-2] + '-O'
return val.upper()
return None

def extract_oficio_salida_formal(subject, body, attachments):
oficio_subj = extract_oficio_tce_exacto(subject)
if oficio_subj:
return oficio_subj

oficio_cuerpo = extract_oficio_tce_exacto(body)
if oficio_cuerpo:
return oficio_cuerpo

body_clean = remove_accents(body)
patterns_leyenda = [
r'(?:se\s+remite|remito|adjunto|hago\s+llegar)\s+(?:el\s+)?oficio\s*(?:nro\.?|n°|num\.?|#)?\s*[:#-]?\s*([A-Za-z0-9_\-\/]{4,35})',
r'oficio\s*(?:nro\.?|n°|num\.?|#)?\s*[:#-]?\s*([A-Za-z0-9_\-\/]{4,35})\s+(?:adjunto|mediante\s+el\s+cual)'
]
for p in patterns_leyenda:
m = re.search(p, body_clean, re.IGNORECASE)
if m:
candidato = m.group(1).strip().strip('.,;')
if "tce" in candidato.lower() and len(candidato) >= 6:
return candidato.upper()

for att in attachments:
oficio_att = extract_oficio_tce_exacto(att)
if oficio_att:
return oficio_att

return "No identificado"

def is_formal_attending_response(subject, body, attachments):
body_norm = remove_accents(body).lower()
formulas_leyenda = [
r'por\s+medio\s+del\s+presente\s+(?:se\s+)?remite',
r'adjunto\s+(?:al\s+presente\s+)?(?:se\s+)?remit[eo]',
r'para\s+(?:los\s+)?fines\s+(?:pertinentes|legales)',
r'en\s+atencion\s+a\s+(?:su\s+)?solicitud',
r'dando\s+atencion\s+a\s+lo\s+solicitado',
r'cumplo\s+con\s+remitir',
r'notific[ao]\s+(?:a\s+usted\s+)?mediante\s+oficio'
]
has_leyenda = any(re.search(pat, body_norm) for pat in formulas_leyenda)
has_oficio_pdf = any(att.lower().endswith('.pdf') for att in attachments)

if extract_oficio_tce_exacto(subject) or extract_oficio_tce_exacto(body):
has_leyenda = True

return has_leyenda, has_oficio_pdf

# PATRONES DE ENTRADA MEJORADOS (Cubre formato estándar y con número al inicio: 0107-CNE-JPEB-2026)
PATRONES_OFICIOS_ENTRADA = [
# 1. Con número al inicio: Ej. 0107-CNE-JPEB-2026 o 107-CNE-DPP-2026
r'\b([0-9]{1,5}-(?:TCE|CNE|JPE[A-Z0-9]*|DP[A-Z0-9]*)(?:-[A-Za-z0-9]+)*-[0-9]{4}(?:-[A-Za-z0-9]+)?)\b',
# 2. Con sigla al inicio: Ej. CNE-JPEB-2026-0107 o CNE-SG-2026-1040
r'\b((?:TCE|CNE|JPE[A-Z0-9]*|DP[A-Z0-9]*)-(?:[A-Za-z0-9]+-)*[0-9]{4}-[0-9]{1,5}(?:-[A-Za-z0-9]+)?)\b',
# 3. Tras mención expresa de oficio/memorando
r'(?:oficio|memorando|memo|ofc)\.?\s*(?:nro\.?|n°|num\.?|#)?\s*([0-9A-Za-z\-_]+(?:-[0-9A-Za-z\-_]+){2,})'
]

def extract_oficio_entrada(text):
for pat in PATRONES_OFICIOS_ENTRADA:
match = re.search(pat, text, re.IGNORECASE)
if match:
val = match.group(1).strip().strip('.,;')
if len(val) >= 6 and not val.lower().startswith(('pdf', 'adjunto', 'solicitud')):
return val.upper()
return "No identificado"

def normalize_subject(subj):
s = remove_accents(subj).lower()
s = re.sub(r'^(re|fwd|rv|enc|fwd:):\s*', '', s, flags=re.IGNORECASE)
return re.sub(r'\s+', ' ', s).strip()

def variants_oficio(ofc_str):
"""Genera variantes con y sin ceros a la izquierda para emparejar con éxito."""
vars_set = {ofc_str.upper()}
# Si tiene número al inicio: 0107-CNE... -> 107-CNE...
m = re.match(r'^0*([0-9]+)(-.+)$', ofc_str, re.IGNORECASE)
if m:
num, resto = m.group(1), m.group(2)
vars_set.add((num + resto).upper())
vars_set.add((num.zfill(4) + resto).upper())
# Si tiene número al final: CNE...-0107 -> CNE...-107
m2 = re.match(r'^(.+-)(0*)([0-9]+)$', ofc_str, re.IGNORECASE)
if m2:
pref, ceros, num = m2.group(1), m2.group(2), m2.group(3)
vars_set.add((pref + num).upper())
vars_set.add((pref + num.zfill(4)).upper())
return vars_set

# ==========================================
# 4. LECTURA PROTEGIDA (SSL + REINTENTOS)
# ==========================================
print(f"[{datetime.now().strftime('%H:%M:%S')}] Conectando a {IMAP_SERVER}:{IMAP_PORT}...")

context = ssl.create_default_context()
context.check_hostname = False
context.verify_mode = ssl.CERT_NONE

mail = None
max_intentos = 3

for intento in range(1, max_intentos + 1):
try:
mail = imaplib.IMAP4_SSL(IMAP_SERVER, IMAP_PORT, ssl_context=context)
mail.login(USERNAME, PASSWORD)
print(f"[{datetime.now().strftime('%H:%M:%S')}] [✓] Autenticación exitosa (intento {intento}).")
break
except (socket.error, imaplib.IMAP4.abort, ssl.SSLError) as e:
print(f"[!] Intento {intento}/{max_intentos} falló ({e}). Reintentando en 5 segundos...")
time.sleep(5)
if intento == max_intentos:
print("[X] Falló la conexión IMAP. Revisa credenciales o espera por bloqueo temporal de IP.")
exit(1)

status, folders = mail.list()
available_folders = [re.search(r'\"([^\"]+)\"$', f.decode('utf-8', errors='ignore')).group(1)
for f in folders if re.search(r'\"([^\"]+)\"$', f.decode('utf-8', errors='ignore'))]

if SENT_FOLDER not in available_folders:
for candidate in ["Enviados", "Elementos enviados", "Sent Messages", "Sent Items", "Sent"]:
found = [fol for fol in available_folders if candidate.lower() in fol.lower()]
if found:
SENT_FOLDER = found[0]
break

def fetch_folder_emails(folder_name, is_inbox=False):
print(f"\n[{datetime.now().strftime('%H:%M:%S')}] Leyendo carpeta '{folder_name}' en MODO SOLO LECTURA...")
mail.select(f'"{folder_name}"', readonly=True)
mail.noop()

query = f'(SINCE "{FECHA_DESDE_STR}")'
status, data = mail.search(None, query)
msg_ids = data[0].split() if status == 'OK' and data[0] else []
print(f"[✓] {len(msg_ids)} correos totales encontrados desde el {FECHA_DESDE_STR}.")

items = []
count_aug = 0
count_sep = 0

for num in msg_ids:
status, msg_data = mail.fetch(num, '(BODY.PEEK[])')
if status != 'OK':
continue

raw = msg_data[0][1]
msg = email.message_from_bytes(raw)

subject = clean_header(msg.get("Subject", ""))
date_tuple = email.utils.parsedate_tz(msg.get("Date"))
msg_date = datetime.fromtimestamp(email.utils.mktime_tz(date_tuple)) if date_tuple else datetime.now()

if msg_date.month == 8:
count_aug += 1
elif msg_date.month == 9:
count_sep += 1

sender = clean_header(msg.get("From", ""))
recipient = clean_header(msg.get("To", ""))
msg_id = msg.get("Message-ID", "").strip()
in_reply_to = msg.get("In-Reply-To", "").strip()
references = msg.get("References", "").strip()

leer_pdf = is_inbox and matches_subject_and_sender_criteria(subject, sender)
body, attachments = parse_email_content_and_attachments(msg, leer_adjuntos_pdf=leer_pdf)
full_content = f"{subject}\n{body}"
oficio_entrada = extract_oficio_entrada(full_content)
oficio_salida = extract_oficio_salida_formal(subject, body, attachments)
has_leyenda, has_oficio_pdf = is_formal_attending_response(subject, body, attachments)

items.append({
"message_id": msg_id,
"in_reply_to": in_reply_to,
"references": references,
"date": msg_date,
"subject": subject,
"norm_subject": normalize_subject(subject),
"sender": sender,
"sender_email": extract_email_address(sender),
"recipient": recipient,
"recipient_email": extract_email_address(recipient),
"oficio_general": oficio_entrada,
"oficio_salida": oficio_salida,
"body": body,
"attachments": attachments,
"has_leyenda": has_leyenda,
"has_oficio_pdf": has_oficio_pdf,
"is_valid_resolution": (has_leyenda and has_oficio_pdf),
"snippet": body[:120].replace('\n', ' ')
})

print(f" -> Agosto: {count_aug} correos | Septiembre: {count_sep} correos")
return items

raw_inbox = fetch_folder_emails(INBOX_FOLDER, is_inbox=True)
raw_sent = fetch_folder_emails(SENT_FOLDER, is_inbox=False)
mail.logout()

inbox_items = [item for item in raw_inbox if matches_subject_and_sender_criteria(item["subject"], item["sender"])]

print(f"\n[✓] Solicitudes válidas filtradas: {len(inbox_items)} de {len(raw_inbox)} recibidos.")
print(f"[✓] Correos en Enviados para cruce: {len(raw_sent)}")

# ==========================================================
# 5. DICCIONARIO DE RESPUESTAS POR OFICIO CON TOLERANCIA
# ==========================================================
respuestas_por_oficio = {}

for s in raw_sent:
if s["oficio_salida"] != "No identificado" and s["has_oficio_pdf"]:
texto_busqueda = f"{s['subject']} {s['body']} {' '.join(s['attachments'])}".upper()
for pat in PATRONES_OFICIOS_ENTRADA:
encontrados = re.findall(pat, texto_busqueda, re.IGNORECASE)
for ofc in encontrados:
for v in variants_oficio(ofc):
if v not in respuestas_por_oficio or s["date"] > respuestas_por_oficio[v]["date"]:
respuestas_por_oficio[v] = s

audit_rows = []
atendidos_count = 0
obs_count = 0

for item in inbox_items:
match_found = None
criterio_match = "Ninguno"
oficio_in = item["oficio_general"]
has_oficio_in = (oficio_in != "No identificado")
# 1. Búsqueda por Oficio en el índice (probando variantes con y sin ceros)
if has_oficio_in:
for v in variants_oficio(oficio_in):
if v in respuestas_por_oficio:
match_found = respuestas_por_oficio[v]
criterio_match = f"N° Oficio Entrada Referenciado ({v})"
break

# 2. Hilo técnico directo (In-Reply-To o References)
if not match_found and item["message_id"]:
for s in raw_sent:
if item["message_id"] in s["in_reply_to"] or item["message_id"] in s["references"]:
if item["sender_email"] and item["sender_email"] in s["recipient_email"]:
match_found = s
criterio_match = "Hilo Directo Técnico (In-Reply-To)"
break

# 3. Búsqueda de texto directo si no estuvo en el índice
if not match_found and has_oficio_in:
for v in variants_oficio(oficio_in):
v_low = v.lower()
for s in raw_sent:
in_body = v_low in s["body"].lower()
in_subj = v_low in s["subject"].lower()
in_att = any(v_low in att.lower() for att in s["attachments"])
if in_body or in_subj or in_att:
match_found = s
criterio_match = f"N° Oficio Entrada Referenciado ({v})"
break
if match_found:
break

# EVALUACIÓN DE ESTADO
if match_found:
fecha_salida = match_found["date"].strftime("%Y-%m-%d %H:%M")
oficio_salida = match_found["oficio_salida"]
dias_resp = round((match_found["date"] - item["date"]).total_seconds() / 86400, 1)
adjuntos_salida = ", ".join(match_found["attachments"]) if match_found["attachments"] else "Sin adjuntos PDF"

tiene_oficio_salida = (oficio_salida != "No identificado")
tiene_pdf = match_found["has_oficio_pdf"]

if has_oficio_in and tiene_oficio_salida and tiene_pdf:
estado = "ATENDIDO"
observacion = f"Contestación formal vinculada por {criterio_match} con oficio {oficio_salida}"
atendidos_count += 1
elif not has_oficio_in:
if tiene_oficio_salida and tiene_pdf:
estado = "ATENDIDO (SIN OFICIO ENTRADA)"
observacion = f"Respuesta formal con {oficio_salida}, pero solicitud no registra oficio formal de entrada"
atendidos_count += 1
else:
estado = "PENDIENTE (SIN OFICIO ENTRADA / VERIFICAR)"
observacion = "Solicitud sin número de oficio de entrada y sin despacho concluyente"
obs_count += 1
else:
estado = "PENDIENTE (FALTA DOC / SUBSANACIÓN)"
motivos = []
if not tiene_oficio_salida:
motivos.append("Sin Oficio formal de salida TCE")
if not tiene_pdf:
motivos.append("Falta adjunto PDF formal")
observacion = "Respuesta intermedia/incompleta: " + "; ".join(motivos)
obs_count += 1
else:
estado = "PENDIENTE"
fecha_salida = ""
oficio_salida = ""
dias_resp = ""
adjuntos_salida = ""
observacion = f"Sin respuesta que referencie el oficio {oficio_in}" if has_oficio_in else "Solicitud sin oficio de entrada y sin respuesta en Enviados"

audit_rows.append({
"Fecha Recepción": item["date"].strftime("%Y-%m-%d %H:%M"),
"Remitente": item["sender"],
"Asunto Recibido": item["subject"],
"N° Oficio Entrada (Solicitud)": item["oficio_general"],
"Estado": estado,
"Fecha Contestación": fecha_salida,
"N° Oficio Salida (Contestación)": oficio_salida,
"Archivos Adjuntos Enviados (PDF)": adjuntos_salida,
"Tiempo Respuesta (Días)": dias_resp,
"Criterio Match": criterio_match,
"Observaciones / Validación": observacion
})

print(f"\n[✓] RESUMEN AL CORTE ({FECHA_AUDITORIA_AHORA.strftime('%Y-%m-%d %H:%M')}):")
print(f" -> ATENDIDOS FORMALES: {atendidos_count}")
print(f" -> PENDIENTES (Falta Doc / Subsanación): {obs_count}")
print(f" -> PENDIENTES SIN RESPUESTA: {len(inbox_items) - atendidos_count - obs_count}")

# ==========================================
# 6. GENERACIÓN DEL EXCEL MULTI-PESTAÑA
# ==========================================
output_file = "auditoria_certificaciones.xlsx"

if not audit_rows:
print("\n[!] AVISO: No se encontraron solicitudes que coincidan con los filtros en el período evaluado.")
audit_rows = [{
"Fecha Recepción": datetime.now().strftime("%Y-%m-%d %H:%M"),
"Remitente": "Sin registros",
"Asunto Recibido": "No se encontraron correos con los criterios de búsqueda",
"N° Oficio Entrada (Solicitud)": "N/A",
"Estado": "SIN DATOS",
"Fecha Contestación": "",
"N° Oficio Salida (Contestación)": "",
"Archivos Adjuntos Enviados (PDF)": "",
"Tiempo Respuesta (Días)": "",
"Criterio Match": "",
"Observaciones / Validación": "Verificar filtros de búsqueda o período"
}]

df_todos = pd.DataFrame(audit_rows)
df_pendientes = df_todos[~df_todos["Estado"].isin(["ATENDIDO", "ATENDIDO (SIN OFICIO ENTRADA)"])].copy()

total_solicitudes = len(df_todos) if df_todos["Estado"].iloc[0] != "SIN DATOS" else 0
atendidos = len(df_todos[df_todos["Estado"] == "ATENDIDO"])
atendidos_sin_oficio = len(df_todos[df_todos["Estado"] == "ATENDIDO (SIN OFICIO ENTRADA)"])
pendientes_doc = len(df_todos[df_todos["Estado"].str.contains("DOC|VERIFICAR", na=False)])
pendientes_cero = len(df_todos[df_todos["Estado"] == "PENDIENTE"])

tiempos_validos = []
for r in audit_rows:
val = r.get("Tiempo Respuesta (Días)")
try:
if val != "" and val is not None:
tiempos_validos.append(float(val))
except (ValueError, TypeError):
pass

tiempo_promedio = round(sum(tiempos_validos) / len(tiempos_validos), 2) if tiempos_validos else 0.0

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
df_todos.to_excel(writer, index=False, sheet_name="Detalle Auditoria", startrow=3)
wb = writer.book
if "Sheet" in wb.sheetnames:
wb.remove(wb["Sheet"])

# 1. Hoja 1: Dashboard Ejecutivo
ws_dash = wb.create_sheet(title="Dashboard Ejecutivo", index=0)

ws_dash.merge_cells("B2:E2")
ws_dash["B2"] = "TRIBUNAL CONTENCIOSO ELECTORAL - PANEL DE CONTROL DE AUDITORÍA"
ws_dash["B2"].font = Font(size=13, bold=True, color="1F497D")
ws_dash["B2"].alignment = Alignment(horizontal="center", vertical="center")

ws_dash.merge_cells("B3:E3")
ws_dash["B3"] = f"Corte al: {FECHA_AUDITORIA_AHORA.strftime('%d/%m/%Y %H:%M')} | Período evaluado: Desde {FECHA_DESDE_STR}"
ws_dash["B3"].font = Font(size=10, italic=True, color="555555")
ws_dash["B3"].alignment = Alignment(horizontal="center", vertical="center")

headers_dash = ["Indicador / Estado", "Total Trámites", "% del Total", "Estado Operativo"]
for col_i, h in enumerate(headers_dash, start=2):
c = ws_dash.cell(row=5, column=col_i, value=h)
c.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
c.font = Font(color="FFFFFF", bold=True)
c.alignment = Alignment(horizontal="center", vertical="center")

rows_metricas = [
("Total Solicitudes Ingresadas", total_solicitudes, "100.0%", "Recepción Oficial", "FFFFFF"),
("Atendidos Formales (Con Oficio Entrada, Salida y PDF)", atendidos, f"{(atendidos/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Despacho Formal Concluido", "D9EAD3"),
("Atendidos (Petición Directa / Sin Oficio Entrada)", atendidos_sin_oficio, f"{(atendidos_sin_oficio/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Resuelto sin Trámite Inicial", "E2EFDA"),
("Pendientes (Falta Doc / Subsanación / Por Verificar)", pendientes_doc, f"{(pendientes_doc/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "En Trámite Intermedio", "FFF2CC"),
("Pendientes Sin Respuesta", pendientes_cero, f"{(pendientes_cero/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Acción Inmediata Requerida", "FCE5CD"),
("Tiempo Promedio de Despacho", f"{tiempo_promedio} días", "-", "Eficiencia Temporal", "EFEFEF")
]

for idx, (label, val, pct, obs, fill_hex) in enumerate(rows_metricas, start=6):
ws_dash.cell(row=idx, column=2, value=label).font = Font(bold=(idx in [6, 11]))
ws_dash.cell(row=idx, column=3, value=val).alignment = Alignment(horizontal="center")
ws_dash.cell(row=idx, column=4, value=pct).alignment = Alignment(horizontal="center")
ws_dash.cell(row=idx, column=5, value=obs).alignment = Alignment(horizontal="center")
for col_c in range(2, 6):
ws_dash.cell(row=idx, column=col_c).fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")

ws_dash.column_dimensions["B"].width = 54
ws_dash.column_dimensions["C"].width = 18
ws_dash.column_dimensions["D"].width = 16
ws_dash.column_dimensions["E"].width = 28

# 2. Hoja 2: Detalle Completo
ws_det = writer.sheets["Detalle Auditoria"]
ws_det.merge_cells("A1:K1")
ws_det["A1"] = "MATRIZ DETALLADA DE SOLICITUDES Y CERTIFICACIONES TCE"
ws_det["A1"].font = Font(size=12, bold=True, color="1F497D")
ws_det["A1"].alignment = Alignment(horizontal="center", vertical="center")

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=10)

for col_idx in range(1, len(df_todos.columns) + 1):
c = ws_det.cell(row=4, column=col_idx)
c.fill = header_fill
c.font = header_font
c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

green_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
light_green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
red_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")

end_row_det = 4 + len(df_todos)
for r in range(5, end_row_det + 1):
st = str(ws_det.cell(row=r, column=5).value)
if st == "ATENDIDO":
f = green_fill
elif "SIN OFICIO ENTRADA" in st and "ATENDIDO" in st:
f = light_green_fill
elif "DOC" in st or "VERIFICAR" in st:
f = yellow_fill
else:
f = red_fill

for c_idx in range(1, len(df_todos.columns) + 1):
ws_det.cell(row=r, column=c_idx).fill = f

for col_idx in range(1, len(df_todos.columns) + 1):
col_let = get_column_letter(col_idx)
max_l = max([len(str(ws_det.cell(row=r, column=col_idx).value or '')) for r in range(4, end_row_det + 1)] or [12])
ws_det.column_dimensions[col_let].width = min(max(max_l + 3, 14), 45)

if len(df_todos) > 0:
ws_det.auto_filter.ref = f"A4:{get_column_letter(len(df_todos.columns))}{end_row_det}"

# 3. Hoja 3: Pendientes y Alertas
if not df_pendientes.empty and total_solicitudes > 0:
df_pendientes.to_excel(writer, index=False, sheet_name="Pendientes y Alertas", startrow=3)
ws_pen = writer.sheets["Pendientes y Alertas"]

ws_pen.merge_cells("A1:K1")
ws_pen["A1"] = "CASOS PENDIENTES DE ATENCIÓN FORMAL (SEGUIMIENTO SECRETARÍA)"
ws_pen["A1"].font = Font(size=12, bold=True, color="C00000")
ws_pen["A1"].alignment = Alignment(horizontal="center", vertical="center")

alert_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
for col_idx in range(1, len(df_pendientes.columns) + 1):
c = ws_pen.cell(row=4, column=col_idx)
c.fill = alert_fill
c.font = Font(color="FFFFFF", bold=True, size=10)
c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

end_row_pen = 4 + len(df_pendientes)
for r in range(5, end_row_pen + 1):
st = str(ws_pen.cell(row=r, column=5).value)
f = yellow_fill if ("DOC" in st or "VERIFICAR" in st) else red_fill
for c_idx in range(1, len(df_pendientes.columns) + 1):
ws_pen.cell(row=r, column=c_idx).fill = f

for col_idx in range(1, len(df_pendientes.columns) + 1):
col_let = get_column_letter(col_idx)
max_l = max([len(str(ws_pen.cell(row=r, column=col_idx).value or '')) for r in range(4, end_row_pen + 1)] or [12])
ws_pen.column_dimensions[col_let].width = min(max(max_l + 3, 14), 45)

ws_pen.auto_filter.ref = f"A4:{get_column_letter(len(df_pendientes.columns))}{end_row_pen}"

wb.active = ws_dash

print(f"\n[✓] Archivo Excel generado con éxito: '{output_file}'.") 
