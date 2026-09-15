"""
Módulo de Auditoría de Certificaciones y Solicitudes TCE
Refactorizado bajo lineamientos de Diseño de Software (UCOM0310 - Ae4).
Preserva el comportamiento funcional exacto de conciliación y reporte Excel.
"""

import email
from email.header import decode_header
import io
import logging
import os
import re
import socket
import ssl
import time
import unicodedata
from datetime import datetime

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

# Silenciar advertencias internas de lectores de PDF
try:
    from pypdf import PdfReader
    PYPDF_DISPONIBLE = True
    logging.getLogger("pypdf").setLevel(logging.ERROR)
except ImportError:
    try:
        from PyPDF2 import PdfReader
        PYPDF_DISPONIBLE = True
        logging.getLogger("PyPDF2").setLevel(logging.ERROR)
    except ImportError:
        PYPDF_DISPONIBLE = False

# ==========================================================
# 1. PARÁMETROS Y CONFIGURACIÓN DEL SISTEMA
# ==========================================================
IMAP_SERVER = os.getenv("TCE_IMAP_SERVER", "mail.tce.gob.ec")
IMAP_PORT = int(os.getenv("TCE_IMAP_PORT", 993))
USERNAME = os.getenv("TCE_IMAP_USER", "secretaria.general@tce.gob.ec")
PASSWORD = os.getenv("TCE_IMAP_PASS", r"Sg.TcE.2025$")

INBOX_FOLDER = "INBOX"
SENT_FOLDER = "Sent"
FECHA_DESDE_STR = "10-Aug-2026"
SOCKET_TIMEOUT = 45

PATRONES_OFICIO_TCE = [
    r'\b(TCE-(?:SG|P|OM)(?:-[A-Za-z0-9]+)*-[0-9]{4}-[0-9]{1,5}-[O0o])\b',
    r'\b(TCE-[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*-[0-9]{4}-[0-9]{1,5}(?:-[O0o])?)\b',
    r'(?:oficio|oficio\s+nro\.?|oficio\s+n°|oficio\s+num\.?)\s*[:#-]?\s*([0-9]{1,5}-[0-9]{4}-TCE(?:-[A-Za-z0-9]+)?)',
    r'(?:oficio|oficio\s+nro\.?|oficio\s+n°|oficio\s+num\.?)\s*[:#-]?\s*(TCE-[A-Za-z0-9_\-\/]+)'
]

PATRONES_OFICIOS_ENTRADA = [
    r'\b([0-9]{1,5}-(?:TCE|CNE|JPE[A-Z0-9]*|DP[A-Z0-9]*)(?:-[A-Za-z0-9]+)*-[0-9]{4}(?:-[A-Za-z0-9]+)?)\b',
    r'\b((?:TCE|CNE|JPE[A-Z0-9]*|DP[A-Z0-9]*)-(?:[A-Za-z0-9]+-)*[0-9]{4}-[0-9]{1,5}(?:-[A-Za-z0-9]+)?)\b',
    r'(?:oficio|memorando|memo|ofc)\.?\s*(?:nro\.?|n°|num\.?|#)?\s*([0-9A-Za-z\-_]+(?:-[0-9A-Za-z\-_]+){2,})'
]

# ==========================================================
# 2. UTILIDADES DE NORMALIZACIÓN Y PARSEO DE MENSAJES
# ==========================================================
def remove_accents(input_str: str) -> str:
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', input_str)
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])

def clean_header(header_val) -> str:
    if not header_val:
        return ""
    decoded = decode_header(header_val)
    parts = []
    for text, enc in decoded:
        if isinstance(text, bytes):
            try:
                parts.append(text.decode(enc or 'utf-8', errors='ignore'))
            except Exception:
                parts.append(text.decode('latin-1', errors='ignore'))
        else:
            parts.append(str(text))
    return "".join(parts).strip()

def extract_email_address(raw_str: str) -> str:
    match = re.search(r'[\w\.-]+@[\w\.-]+', raw_str)
    return match.group(0).lower() if match else ""

def normalize_subject(subj: str) -> str:
    s = remove_accents(subj).lower()
    s = re.sub(r'^(re|fwd|rv|enc|fwd:):\s*', '', s, flags=re.IGNORECASE)
    return re.sub(r'\s+', ' ', s).strip()

def matches_subject_and_sender_criteria(subject: str, sender_raw: str) -> bool:
    sender_clean = remove_accents(sender_raw).lower()
    sender_email = extract_email_address(sender_raw)
    text_clean = remove_accents(subject).lower()
    text_raw = subject.lower()
    
    if "documentos@cne.gob.ec" in sender_email or "edoc" in sender_clean or "edoc" in text_clean:
        return False
    
    strict_patterns = [
        r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+(?:presentados?|interpuestos?)\s+(?:en\s+contra\s+de|contra)\s+(?:las\s+)?resoluci[oó]n(?:es)?',
        r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+(?:presentados?|interpuestos?)\s+(?:en\s+contra\s+de|contra)',
        r'(?:recursos?|resoluci[oó]n(?:es)?)\s+(?:en\s+contra\s+de|contra)\s+(?:las\s+)?resoluci[oó]n(?:es)?\s+(?:al|ante\s+el)\s+tce',
        r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+subjetivos?(?:\s+contenciosos?)?(?:\s+electorales?)?',
        r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+subjetivos?',
        r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+pendientes?',
        r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?\s+pendientes?',
        r'(?:solicitud|solicito|peticion|pedido)?\s*(?:de\s+)?(?:no\s+)?(?:tener|poseer|registrar|existir)\s+recursos?\s+pendientes?',
        r'no\s+tener\s+recursos?\s+pendientes?',
        r'(?:solicitud|solicito|peticion|pedido)\s+(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:las\s+)?resoluci[oó]n(?:es)?',
        r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:las\s+)?resoluci[oó]n(?:es)?',
        r'(?:solicitud|solicito|peticion|pedido)\s+(?:de\s+)?certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?',
        r'certificaci[oó]n\s+(?:de\s+|sobre\s+|del?\s+)?(?:los\s+)?recursos?',
        r'certificaci[oó]n\s+(?:de\s+)?no\s+(?:haber\s+)?interpuesto\s+recursos?',
        r'certificaci[oó]n',
        r'certificaciones',
        r'solicitud\s+(?:de\s+)?certificaci[oó]n',
        r'solicito\s+certificaci[oó]n',
        r'certificacion',
        r'certificación',
        r'junta[s]?\s+provincial(?:es)?',
        r'junta[s]?\s+provincial(?:es)?\s+electoral(?:es)?',
        r'\bjpe[a-z0-9_\-]*\b',
        r'delegaci[oó]n\s+provincial',
        r'interpuest[oa]s?',
        r'interposici[oó]n',
        r'interposicion',
        r'interposición'
    ]
    return any(re.search(pat, text_clean) or re.search(pat, text_raw) for pat in strict_patterns)

def parse_email_content_and_attachments(msg, leer_adjuntos_pdf=False):
    body = ""
    attachments = []
    pdf_text_content = ""
    
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            cdispo = str(part.get('Content-Disposition', ''))
            filename = part.get_filename()
            
            if filename:
                clean_fname = clean_header(filename)
                fname_lower = clean_fname.lower()
                if not (fname_lower.endswith('.doc') or fname_lower.endswith('.docx')):
                    attachments.append(clean_fname)
                
                if leer_adjuntos_pdf and fname_lower.endswith('.pdf') and PYPDF_DISPONIBLE:
                    try:
                        payload = part.get_payload(decode=True)
                        if payload:
                            pdf_file = io.BytesIO(payload)
                            reader = PdfReader(pdf_file, strict=False)
                            for page_idx in range(min(2, len(reader.pages))):
                                extracted = reader.pages[page_idx].extract_text()
                                if extracted:
                                    pdf_text_content += extracted + " "
                    except Exception:
                        pass
            elif 'attachment' in cdispo:
                attachments.append("adjunto_sin_nombre")
                
            if ctype in ['text/plain', 'text/html'] and 'attachment' not in cdispo:
                try:
                    payload = part.get_payload(decode=True)
                    if payload:
                        decoded = payload.decode('utf-8', errors='ignore')
                        decoded = re.sub(r'<[^>]+>', ' ', decoded)
                        body += decoded + " "
                except Exception:
                    pass
    else:
        try:
            payload = msg.get_payload(decode=True)
            if payload:
                body = payload.decode('utf-8', errors='ignore')
        except Exception:
            pass

    full_body = (body + " " + pdf_text_content).strip()
    return full_body, attachments

# ==========================================================
# 3. EXTRACCIÓN Y RECONOCIMIENTO DE OFICIOS
# ==========================================================
def extract_oficio_tce_exacto(text: str):
    if not text:
        return None
    for pat in PATRONES_OFICIO_TCE:
        match = re.search(pat, text, re.IGNORECASE)
        if match:
            val = match.group(1).strip().strip('.,;')
            if val.endswith('-0'):
                val = val[:-2] + '-O'
            return val.upper()
    return None

def extract_oficio_salida_formal(subject: str, body: str, attachments: list) -> str:
    oficio = extract_oficio_tce_exacto(subject) or extract_oficio_tce_exacto(body)
    if oficio:
        return oficio

    body_clean = remove_accents(body)
    patterns_leyenda = [
        r'(?:se\s+remite|remito|adjunto|hago\s+llegar)\s+(?:el\s+)?oficio\s*(?:nro\.?|n°|num\.?|#)?\s*[:#-]?\s*([A-Za-z0-9_\-\/]{4,35})',
        r'oficio\s*(?:nro\.?|n°|num\.?|#)?\s*[:#-]?\s*([A-Za-z0-9_\-\/]{4,35})\s+(?:adjunto|mediante\s+el\s+cual)'
    ]
    for p in patterns_leyenda:
        m = re.search(p, body_clean, re.IGNORECASE)
        if m:
            candidato = m.group(1).strip().strip('.,;')
            if "tce" in candidato.lower() and len(candidato) >= 6:
                return candidato.upper()

    for att in attachments:
        oficio_att = extract_oficio_tce_exacto(att)
        if oficio_att:
            return oficio_att

    return "No identificado"

def is_formal_attending_response(subject: str, body: str, attachments: list):
    body_norm = remove_accents(body).lower()
    formulas_leyenda = [
        r'por\s+medio\s+del\s+presente\s+(?:se\s+)?remite',
        r'adjunto\s+(?:al\s+presente\s+)?(?:se\s+)?remit[eo]',
        r'para\s+(?:los\s+)?fines\s+(?:pertinentes|legales)',
        r'en\s+atencion\s+a\s+(?:su\s+)?solicitud',
        r'dando\s+atencion\s+a\s+lo\s+solicitado',
        r'cumplo\s+con\s+remitir',
        r'notific[ao]\s+(?:a\s+usted\s+)?mediante\s+oficio'
    ]
    has_leyenda = any(re.search(pat, body_norm) for pat in formulas_leyenda)
    has_oficio_pdf = any(att.lower().endswith('.pdf') for att in attachments)

    if extract_oficio_tce_exacto(subject) or extract_oficio_tce_exacto(body):
        has_leyenda = True

    return has_leyenda, has_oficio_pdf

def extract_oficio_entrada(text: str) -> str:
    for pat in PATRONES_OFICIOS_ENTRADA:
        match = re.search(pat, text, re.IGNORECASE)
        if match:
            val = match.group(1).strip().strip('.,;')
            if len(val) >= 6 and not val.lower().startswith(('pdf', 'adjunto', 'solicitud')):
                return val.upper()
    return "No identificado"

def variants_oficio(ofc_str: str) -> set:
    vars_set = {ofc_str.upper()}
    m = re.match(r'^0*([0-9]+)(-.+)$', ofc_str, re.IGNORECASE)
    if m:
        num, resto = m.group(1), m.group(2)
        vars_set.add((num + resto).upper())
        vars_set.add((num.zfill(4) + resto).upper())
    m2 = re.match(r'^(.+-)(0*)([0-9]+)$', ofc_str, re.IGNORECASE)
    if m2:
        pref, _, num = m2.group(1), m2.group(2), m2.group(3)
        vars_set.add((pref + num).upper())
        vars_set.add((pref + num.zfill(4)).upper())
    return vars_set

# ==========================================================
# 4. CONEXIÓN Y DESCARGA IMAP
# ==========================================================
def conectar_imap(server: str, port: int, user: str, password: str, max_retries: int = 3):
    import imaplib
    socket.setdefaulttimeout(SOCKET_TIMEOUT)
    context = ssl.create_default_context()
    context.check_hostname = False
    context.verify_mode = ssl.CERT_NONE

    for intento in range(1, max_retries + 1):
        try:
            mail = imaplib.IMAP4_SSL(server, port, ssl_context=context)
            mail.login(user, password)
            print(f"[{datetime.now().strftime('%H:%M:%S')}] [✓] Conectado e identificado exitosamente.")
            return mail
        except (socket.error, imaplib.IMAP4.abort, ssl.SSLError) as e:
            print(f"[!] Reintento {intento}/{max_retries} por error: {e}")
            time.sleep(5)
    raise ConnectionError("No se pudo conectar al servidor IMAP tras múltiples intentos.")

def resolver_nombre_carpeta_enviados(mail, fallback_folder: str = "Sent") -> str:
    _, folders = mail.list()
    available_folders = [
        re.search(r'\"([^\"]+)\"$', f.decode('utf-8', errors='ignore')).group(1) 
        for f in folders if re.search(r'\"([^\"]+)\"$', f.decode('utf-8', errors='ignore'))
    ]
    if fallback_folder in available_folders:
        return fallback_folder
    for candidate in ["Enviados", "Elementos enviados", "Sent Messages", "Sent Items", "Sent"]:
        for fol in available_folders:
            if candidate.lower() in fol.lower():
                return fol
    return fallback_folder

def fetch_folder_emails(mail, folder_name: str, fecha_desde: str, is_inbox: bool = False):
    print(f"\n[{datetime.now().strftime('%H:%M:%S')}] Leyendo carpeta '{folder_name}' en solo lectura...")
    mail.select(f'"{folder_name}"', readonly=True)
    mail.noop()

    status, data = mail.search(None, f'(SINCE "{fecha_desde}")')
    msg_ids = data[0].split() if status == 'OK' and data[0] else []
    print(f"[✓] {len(msg_ids)} correos encontrados desde {fecha_desde}.")

    items = []
    count_aug, count_sep = 0, 0

    for num in msg_ids:
        status, msg_data = mail.fetch(num, '(BODY.PEEK[])')
        if status != 'OK':
            continue

        raw = msg_data[0][1]
        msg = email.message_from_bytes(raw)

        subject = clean_header(msg.get("Subject", ""))
        date_tuple = email.utils.parsedate_tz(msg.get("Date"))
        msg_date = datetime.fromtimestamp(email.utils.mktime_tz(date_tuple)) if date_tuple else datetime.now()

        if msg_date.month == 8:
            count_aug += 1
        elif msg_date.month == 9:
            count_sep += 1

        sender = clean_header(msg.get("From", ""))
        recipient = clean_header(msg.get("To", ""))
        leer_pdf = is_inbox and matches_subject_and_sender_criteria(subject, sender)
        body, attachments = parse_email_content_and_attachments(msg, leer_adjuntos_pdf=leer_pdf)
        
        has_leyenda, has_oficio_pdf = is_formal_attending_response(subject, body, attachments)

        items.append({
            "message_id": msg.get("Message-ID", "").strip(),
            "in_reply_to": msg.get("In-Reply-To", "").strip(),
            "references": msg.get("References", "").strip(),
            "date": msg_date,
            "subject": subject,
            "norm_subject": normalize_subject(subject),
            "sender": sender,
            "sender_email": extract_email_address(sender),
            "recipient": recipient,
            "recipient_email": extract_email_address(recipient),
            "oficio_general": extract_oficio_entrada(f"{subject}\n{body}"),
            "oficio_salida": extract_oficio_salida_formal(subject, body, attachments),
            "body": body,
            "attachments": attachments,
            "has_leyenda": has_leyenda,
            "has_oficio_pdf": has_oficio_pdf,
            "is_valid_resolution": (has_leyenda and has_oficio_pdf)
        })

    print(f"    -> Agosto: {count_aug} | Septiembre: {count_sep}")
    return items

# ==========================================================
# 5. CONCILIACIÓN DE OFICIOS Y LÓGICA DE AUDITORÍA
# ==========================================================
def indexar_respuestas_por_oficio(raw_sent: list) -> dict:
    respuestas = {}
    for s in raw_sent:
        if s["oficio_salida"] != "No identificado" and s["has_oficio_pdf"]:
            texto_busqueda = f"{s['subject']} {s['body']} {' '.join(s['attachments'])}".upper()
            for pat in PATRONES_OFICIOS_ENTRADA:
                for ofc in re.findall(pat, texto_busqueda, re.IGNORECASE):
                    for v in variants_oficio(ofc):
                        if v not in respuestas or s["date"] > respuestas[v]["date"]:
                            respuestas[v] = s
    return respuestas

def buscar_respuesta_para_solicitud(solicitud: dict, respuestas_por_oficio: dict, raw_sent: list):
    oficio_in = solicitud["oficio_general"]
    has_oficio_in = (oficio_in != "No identificado")

    # Estrategia 1: Match en índice de oficios
    if has_oficio_in:
        for v in variants_oficio(oficio_in):
            if v in respuestas_por_oficio:
                return respuestas_por_oficio[v], f"N° Oficio Entrada Referenciado ({v})"

    # Estrategia 2: Hilo técnico directo de correo
    if solicitud["message_id"]:
        for s in raw_sent:
            if (solicitud["message_id"] in s["in_reply_to"] or solicitud["message_id"] in s["references"]):
                if solicitud["sender_email"] and solicitud["sender_email"] in s["recipient_email"]:
                    return s, "Hilo Directo Técnico (In-Reply-To)"

    # Estrategia 3: Búsqueda textual completa en contenido de respuesta
    if has_oficio_in:
        for v in variants_oficio(oficio_in):
            v_low = v.lower()
            for s in raw_sent:
                if (v_low in s["body"].lower() or v_low in s["subject"].lower() or any(v_low in att.lower() for att in s["attachments"])):
                    return s, f"N° Oficio Entrada Referenciado ({v})"

    return None, "Ninguno"

def evaluar_estado_tramite(item: dict, match_found: dict, criterio_match: str) -> tuple:
    has_oficio_in = (item["oficio_general"] != "No identificado")

    if not match_found:
        obs = f"Sin respuesta que referencie el oficio {item['oficio_general']}" if has_oficio_in else "Solicitud sin oficio de entrada y sin respuesta en Enviados"
        return "PENDIENTE", "", "", "", "", obs

    fecha_salida = match_found["date"].strftime("%Y-%m-%d %H:%M")
    oficio_salida = match_found["oficio_salida"]
    dias_resp = round((match_found["date"] - item["date"]).total_seconds() / 86400, 1)
    adjuntos_salida = ", ".join(match_found["attachments"]) if match_found["attachments"] else "Sin adjuntos PDF"

    tiene_oficio_salida = (oficio_salida != "No identificado")
    tiene_pdf = match_found["has_oficio_pdf"]

    if has_oficio_in and tiene_oficio_salida and tiene_pdf:
        return "ATENDIDO", fecha_salida, oficio_salida, dias_resp, adjuntos_salida, f"Contestación formal vinculada por {criterio_match} con oficio {oficio_salida}"

    if not has_oficio_in:
        if tiene_oficio_salida and tiene_pdf:
            return "ATENDIDO (SIN OFICIO ENTRADA)", fecha_salida, oficio_salida, dias_resp, adjuntos_salida, f"Respuesta formal con {oficio_salida}, pero solicitud no registra oficio formal de entrada"
        return "PENDIENTE (SIN OFICIO ENTRADA / VERIFICAR)", fecha_salida, oficio_salida, dias_resp, adjuntos_salida, "Solicitud sin número de oficio de entrada y sin despacho concluyente"

    motivos = []
    if not tiene_oficio_salida:
        motivos.append("Sin Oficio formal de salida TCE")
    if not tiene_pdf:
        motivos.append("Falta adjunto PDF formal")
    return "PENDIENTE (FALTA DOC / SUBSANACIÓN)", fecha_salida, oficio_salida, dias_resp, adjuntos_salida, "Respuesta intermedia/incompleta: " + "; ".join(motivos)

def conciliar_auditoria(inbox_items: list, raw_sent: list) -> tuple:
    indice_respuestas = indexar_respuestas_por_oficio(raw_sent)
    filas = []
    total_atendidos = 0
    total_observaciones = 0

    for item in inbox_items:
        match_found, criterio = buscar_respuesta_para_solicitud(item, indice_respuestas, raw_sent)
        estado, f_salida, of_salida, dias, adjuntos, observacion = evaluar_estado_tramite(item, match_found, criterio)

        if "ATENDIDO" in estado:
            total_atendidos += 1
        elif "PENDIENTE" in estado and ("DOC" in estado or "VERIFICAR" in estado):
            total_observaciones += 1

        filas.append({
            "Fecha Recepción": item["date"].strftime("%Y-%m-%d %H:%M"),
            "Remitente": item["sender"],
            "Asunto Recibido": item["subject"],
            "N° Oficio Entrada (Solicitud)": item["oficio_general"],
            "Estado": estado,
            "Fecha Contestación": f_salida,
            "N° Oficio Salida (Contestación)": of_salida,
            "Archivos Adjuntos Enviados (PDF)": adjuntos,
            "Tiempo Respuesta (Días)": dias,
            "Criterio Match": criterio,
            "Observaciones / Validación": observacion
        })

    return filas, total_atendidos, total_observaciones

# ==========================================================
# 6. EXPORTACIÓN Y FORMATEO DE DASHBOARD EN EXCEL
# ==========================================================
def generar_reporte_excel(filas_audit: list, fecha_corte: datetime, output_file: str):
    if not filas_audit:
        filas_audit = [{
            "Fecha Recepción": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "Remitente": "Sin registros",
            "Asunto Recibido": "No se encontraron correos con los criterios de búsqueda",
            "N° Oficio Entrada (Solicitud)": "N/A",
            "Estado": "SIN DATOS",
            "Fecha Contestación": "",
            "N° Oficio Salida (Contestación)": "",
            "Archivos Adjuntos Enviados (PDF)": "",
            "Tiempo Respuesta (Días)": "",
            "Criterio Match": "",
            "Observaciones / Validación": "Verificar filtros de búsqueda o período"
        }]

    df_todos = pd.DataFrame(filas_audit)
    df_pendientes = df_todos[~df_todos["Estado"].isin(["ATENDIDO", "ATENDIDO (SIN OFICIO ENTRADA)"])].copy()

    total_solicitudes = len(df_todos) if df_todos["Estado"].iloc[0] != "SIN DATOS" else 0
    atendidos = len(df_todos[df_todos["Estado"] == "ATENDIDO"])
    atendidos_sin_oficio = len(df_todos[df_todos["Estado"] == "ATENDIDO (SIN OFICIO ENTRADA)"])
    pendientes_doc = len(df_todos[df_todos["Estado"].str.contains("DOC|VERIFICAR", na=False)])
    pendientes_cero = len(df_todos[df_todos["Estado"] == "PENDIENTE"])

    tiempos = [float(r["Tiempo Respuesta (Días)"]) for r in filas_audit if r.get("Tiempo Respuesta (Días)") not in ["", None]]
    tiempo_promedio = round(sum(tiempos) / len(tiempos), 2) if tiempos else 0.0

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_todos.to_excel(writer, index=False, sheet_name="Detalle Auditoria", startrow=3)
        wb = writer.book
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        # Pestaña 1: Dashboard Ejecutivo
        ws_dash = wb.create_sheet(title="Dashboard Ejecutivo", index=0)
        ws_dash.merge_cells("B2:E2")
        ws_dash["B2"] = "TRIBUNAL CONTENCIOSO ELECTORAL - PANEL DE CONTROL DE AUDITORÍA"
        ws_dash["B2"].font = Font(size=13, bold=True, color="1F497D")
        ws_dash["B2"].alignment = Alignment(horizontal="center", vertical="center")

        ws_dash.merge_cells("B3:E3")
        ws_dash["B3"] = f"Corte al: {fecha_corte.strftime('%d/%m/%Y %H:%M')} | Período evaluado: Desde {FECHA_DESDE_STR}"
        ws_dash["B3"].font = Font(size=10, italic=True, color="555555")
        ws_dash["B3"].alignment = Alignment(horizontal="center", vertical="center")

        headers_dash = ["Indicador / Estado", "Total Trámites", "% del Total", "Estado Operativo"]
        for col_i, h in enumerate(headers_dash, start=2):
            c = ws_dash.cell(row=5, column=col_i, value=h)
            c.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")

        metricas = [
            ("Total Solicitudes Ingresadas", total_solicitudes, "100.0%", "Recepción Oficial", "FFFFFF"),
            ("Atendidos Formales (Con Oficio Entrada, Salida y PDF)", atendidos, f"{(atendidos/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Despacho Formal Concluido", "D9EAD3"),
            ("Atendidos (Petición Directa / Sin Oficio Entrada)", atendidos_sin_oficio, f"{(atendidos_sin_oficio/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Resuelto sin Trámite Inicial", "E2EFDA"),
            ("Pendientes (Falta Doc / Subsanación / Por Verificar)", pendientes_doc, f"{(pendientes_doc/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "En Trámite Intermedio", "FFF2CC"),
            ("Pendientes Sin Respuesta", pendientes_cero, f"{(pendientes_cero/total_solicitudes*100):.1f}%" if total_solicitudes else "0%", "Acción Inmediata Requerida", "FCE5CD"),
            ("Tiempo Promedio de Despacho", f"{tiempo_promedio} días", "-", "Eficiencia Temporal", "EFEFEF")
        ]

        for idx, (lbl, val, pct, obs, color) in enumerate(metricas, start=6):
            ws_dash.cell(row=idx, column=2, value=lbl).font = Font(bold=(idx in [6, 11]))
            ws_dash.cell(row=idx, column=3, value=val).alignment = Alignment(horizontal="center")
            ws_dash.cell(row=idx, column=4, value=pct).alignment = Alignment(horizontal="center")
            ws_dash.cell(row=idx, column=5, value=obs).alignment = Alignment(horizontal="center")
            for col_c in range(2, 6):
                ws_dash.cell(row=idx, column=col_c).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        ws_dash.column_dimensions["B"].width = 54
        ws_dash.column_dimensions["C"].width = 18
        ws_dash.column_dimensions["D"].width = 16
        ws_dash.column_dimensions["E"].width = 28

        # Pestaña 2: Detalle
        ws_det = writer.sheets["Detalle Auditoria"]
        ws_det.merge_cells("A1:K1")
        ws_det["A1"] = "MATRIZ DETALLADA DE SOLICITUDES Y CERTIFICACIONES TCE"
        ws_det["A1"].font = Font(size=12, bold=True, color="1F497D")
        ws_det["A1"].alignment = Alignment(horizontal="center", vertical="center")

        estilos = {
            "header": PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid"),
            "green": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
            "light_green": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            "yellow": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            "red": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
        }

        for col_idx in range(1, len(df_todos.columns) + 1):
            c = ws_det.cell(row=4, column=col_idx)
            c.fill = estilos["header"]
            c.font = Font(color="FFFFFF", bold=True, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        end_row_det = 4 + len(df_todos)
        for r in range(5, end_row_det + 1):
            st = str(ws_det.cell(row=r, column=5).value)
            fill_actual = estilos["green"] if st == "ATENDIDO" else estilos["light_green"] if ("SIN OFICIO ENTRADA" in st and "ATENDIDO" in st) else estilos["yellow"] if ("DOC" in st or "VERIFICAR" in st) else estilos["red"]
            for c_idx in range(1, len(df_todos.columns) + 1):
                ws_det.cell(row=r, column=c_idx).fill = fill_actual

        for col_idx in range(1, len(df_todos.columns) + 1):
            col_let = get_column_letter(col_idx)
            max_l = max([len(str(ws_det.cell(row=r, column=col_idx).value or '')) for r in range(4, end_row_det + 1)] or [12])
            ws_det.column_dimensions[col_let].width = min(max(max_l + 3, 14), 45)

        if len(df_todos) > 0:
            ws_det.auto_filter.ref = f"A4:{get_column_letter(len(df_todos.columns))}{end_row_det}"

        # Pestaña 3: Pendientes y Alertas
        if not df_pendientes.empty and total_solicitudes > 0:
            df_pendientes.to_excel(writer, index=False, sheet_name="Pendientes y Alertas", startrow=3)
            ws_pen = writer.sheets["Pendientes y Alertas"]
            ws_pen.merge_cells("A1:K1")
            ws_pen["A1"] = "CASOS PENDIENTES DE ATENCIÓN FORMAL (SEGUIMIENTO SECRETARÍA)"
            ws_pen["A1"].font = Font(size=12, bold=True, color="C00000")
            ws_pen["A1"].alignment = Alignment(horizontal="center", vertical="center")

            alert_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            for col_idx in range(1, len(df_pendientes.columns) + 1):
                c = ws_pen.cell(row=4, column=col_idx)
                c.fill = alert_fill
                c.font = Font(color="FFFFFF", bold=True, size=10)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            end_row_pen = 4 + len(df_pendientes)
            for r in range(5, end_row_pen + 1):
                st = str(ws_pen.cell(row=r, column=5).value)
                f_pen = estilos["yellow"] if ("DOC" in st or "VERIFICAR" in st) else estilos["red"]
                for c_idx in range(1, len(df_pendientes.columns) + 1):
                    ws_pen.cell(row=r, column=c_idx).fill = f_pen

            for col_idx in range(1, len(df_pendientes.columns) + 1):
                col_let = get_column_letter(col_idx)
                max_l = max([len(str(ws_pen.cell(row=r, column=col_idx).value or '')) for r in range(4, end_row_pen + 1)] or [12])
                ws_pen.column_dimensions[col_let].width = min(max(max_l + 3, 14), 45)

            ws_pen.auto_filter.ref = f"A4:{get_column_letter(len(df_pendientes.columns))}{end_row_pen}"

        wb.active = ws_dash

# ==========================================================
# 7. FUNCIÓN PRINCIPAL / ORQUESTADOR
# ==========================================================
def main():
    fecha_auditoria = datetime.now()
    print("=" * 65)
    print(" AUDITORÍA DE CERTIFICACIONES Y SOLICITUDES TCE (REFACTORIZADO)")
    print(f" Fecha y Hora de Ejecución: {fecha_auditoria.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f" Período Evaluado: Desde {FECHA_DESDE_STR} hasta {fecha_auditoria.strftime('%d-%b-%Y')}")
    print("=" * 65)

    mail_client = conectar_imap(IMAP_SERVER, IMAP_PORT, USERNAME, PASSWORD)
    carpeta_sent = resolver_nombre_carpeta_enviados(mail_client, SENT_FOLDER)

    raw_inbox = fetch_folder_emails(mail_client, INBOX_FOLDER, FECHA_DESDE_STR, is_inbox=True)
    raw_sent = fetch_folder_emails(mail_client, carpeta_sent, FECHA_DESDE_STR, is_inbox=False)
    mail_client.logout()

    inbox_items = [item for item in raw_inbox if matches_subject_and_sender_criteria(item["subject"], item["sender"])]
    print(f"\n[✓] Solicitudes válidas filtradas: {len(inbox_items)} de {len(raw_inbox)} recibidos.")
    print(f"[✓] Correos en Enviados para cruce: {len(raw_sent)}")

    audit_rows, atendidos, pendientes_obs = conciliar_auditoria(inbox_items, raw_sent)

    print(f"\n[✓] RESUMEN AL CORTE ({fecha_auditoria.strftime('%Y-%m-%d %H:%M')}):")
    print(f"   -> ATENDIDOS FORMALES: {atendidos}")
    print(f"   -> PENDIENTES (Falta Doc / Subsanación): {pendientes_obs}")
    print(f"   -> PENDIENTES SIN RESPUESTA: {len(inbox_items) - atendidos - pendientes_obs}")

    output_file = "auditoria_certificaciones.xlsx"
    generar_reporte_excel(audit_rows, fecha_auditoria, output_file)
    print(f"\n[✓] Archivo Excel generado con éxito: '{output_file}'.")

if __name__ == "__main__":
    main()